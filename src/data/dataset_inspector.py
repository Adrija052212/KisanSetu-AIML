from pathlib import Path
from openpyxl import load_workbook


def inspect_excel_dataset(file_path):
    """
    Efficiently inspect a large Excel dataset without
    loading the entire workbook into a pandas DataFrame.
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Dataset not found: {file_path}"
        )

    print("\nOpening Excel workbook...")
    print("-" * 50)

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True
    )

    print("Workbook opened successfully.")

    print("\nSheets")
    print("-" * 50)

    print(workbook.sheetnames)

    for sheet_name in workbook.sheetnames:

        sheet = workbook[sheet_name]

        print(f"\nSheet: {sheet_name}")
        print(f"Rows: {sheet.max_row:,}")
        print(f"Columns: {sheet.max_column:,}")

        # Read only the first row
        rows = sheet.iter_rows(
            min_row=1,
            max_row=1,
            values_only=True
        )

        headers = list(next(rows))

        print("\nColumns:")
        for column in headers:
            print(f"  - {column}")

        # Read first 5 data rows
        print("\nFirst 5 rows:")

        rows = sheet.iter_rows(
            min_row=2,
            max_row=min(6, sheet.max_row),
            values_only=True
        )

        for row in rows:
            print(row)

    workbook.close()

    print("\nInspection completed.")


if __name__ == "__main__":

    file_path = (
        "data/raw/SIH_Cleaned_Dataset.xlsx"
    )

    inspect_excel_dataset(file_path)